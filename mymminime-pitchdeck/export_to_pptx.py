import os
import asyncio
from playwright.async_api import async_playwright
from pptx import Presentation
from pptx.util import Inches

SLIDES = [
    "01-cover.html", "02-thesis.html", "03-empathize.html", "04-define.html",
    "05-causal.html", "06-reframe.html", "07-solution.html", "08-prototype.html",
    "09-ladder.html", "10-gtm.html", "11-bm.html", "12-test.html",
    "13-moat.html", "14-os.html", "15-benchmark.html", "16-fundraising.html",
    "17-milestones.html", "18-closing.html"
]

async def capture_and_build():
    print("Starting Playwright...")
    async with async_playwright() as p:
        browser = await p.chromium.launch()
        page = await browser.new_page(viewport={'width': 1920, 'height': 1080})
        
        prs = Presentation()
        # Set slide size to 16:9
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        blank_slide_layout = prs.slide_layouts[6]
        
        base_dir = os.path.abspath("slides")
        os.makedirs(".tmp_shots", exist_ok=True)
        
        for slide_file in SLIDES:
            file_url = f"file://{os.path.join(base_dir, slide_file)}"
            print(f"Capturing: {slide_file}")
            await page.goto(file_url, wait_until="networkidle")
            # small delay for fonts
            await page.wait_for_timeout(1000)
            
            shot_path = f".tmp_shots/{slide_file}.png"
            await page.screenshot(path=shot_path)
            
            slide = prs.slides.add_slide(blank_slide_layout)
            slide.shapes.add_picture(shot_path, 0, 0, width=prs.slide_width, height=prs.slide_height)
            
        await browser.close()
        
        output_name = "MyMiniMe_PitchDeck.pptx"
        prs.save(output_name)
        print(f"Successfully generated: {output_name}")

asyncio.run(capture_and_build())
