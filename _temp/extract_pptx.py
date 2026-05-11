
import pptx
import json
import os

pptx_path = "/Users/andy/Antigravity/2026/5월/huashu design/mymminime_design_thinking_pitchdeck_45억.pptx"
output_path = "/Users/andy/Antigravity/2026/5월/huashu design/_temp/pptx_content.json"

os.makedirs(os.path.dirname(output_path), exist_ok=True)

prs = pptx.Presentation(pptx_path)
slides_content = []

for i, slide in enumerate(prs.slides):
    slide_data = {
        "slide_number": i + 1,
        "title": "",
        "text_blocks": []
    }
    
    if slide.shapes.title:
        slide_data["title"] = slide.shapes.title.text
        
    for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text and shape != slide.shapes.title:
            slide_data["text_blocks"].append(shape.text)
            
    slides_content.append(slide_data)

with open(output_path, "w", encoding="utf-8") as f:
    json.dump(slides_content, f, ensure_ascii=False, indent=2)

print(f"Extracted {len(slides_content)} slides to {output_path}")
